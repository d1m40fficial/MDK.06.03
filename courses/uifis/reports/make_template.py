"""Generate the Printable report template for Practice 1 (УФИС).

Run:  python make_template.py
Writes: УФИС_Практика-1_Шаблон-отчёта.docx  (next to this file)
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm

OUT = Path(__file__).with_name("УФИС_Практика-1_Шаблон-отчёта.docx")
QUESTIONS = [
    "Почему признаки классификации должны быть независимыми?",
    "Может ли одна система быть пакетной и оперативной одновременно?",
    "Чем масштаб отличается от сферы применения?",
]
ROWS = [
    ("1", "Банковская АБС"),
    ("2", "1С:Бухгалтерия малого предприятия"),
    ("3", "Система диспетчеризации завода"),
    ("4", "Интернет-магазин"),
    ("5", "Медицинская картотека поликлиники"),
]


def line(doc, text, bold=False, size=11, align=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    if align is not None:
        p.alignment = align
    return p


def slot(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(11)


def main():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)

    line(doc, "Устройство и функционирование информационной системы (МДК.06.03)", bold=True, size=14,
         align=WD_ALIGN_PARAGRAPH.CENTER)
    line(doc, "Отчёт по практической работе №1", bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER)
    line(doc, "Тема: Классификация информационных систем", align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    line(doc, "ФИО студента: ______________________________________")
    line(doc, "Группа: ______________________  Дата: _______________")
    doc.add_paragraph()

    line(doc, "Цель работы", bold=True)
    line(doc, "________________________________________________________________________")
    line(doc, "________________________________________________________________________")
    doc.add_paragraph()

    line(doc, "Краткие теоретические сведения", bold=True)
    line(doc, "Перечислите четыре признака классификации ИС и приведите по 2–3 значения каждого:")
    line(doc, "1. Сфера применения: ____________________________________________________")
    line(doc, "2. Масштаб: _____________________________________________________________")
    line(doc, "3. Способ обработки: ____________________________________________________")
    line(doc, "4. Архитектура: _________________________________________________________")
    doc.add_paragraph()

    line(doc, "Таблица результатов классификации", bold=True)
    cols = ["№", "Информационная система", "Сфера", "Масштаб", "Способ обработки", "Архитектура",
            "Обоснование"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
    for n, name in ROWS:
        cells = t.add_row().cells
        cells[0].text = n
        cells[1].text = name
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
    doc.add_paragraph()

    line(doc, "Скриншоты", bold=True)
    slot(doc, "[ Место для скриншота 1: шапка отчёта с заполненными полями ]")
    doc.add_paragraph()
    slot(doc, "[ Место для скриншота 2: карточка «Банковская АБС» в классификаторе ]")
    doc.add_paragraph()

    line(doc, "Вывод", bold=True)
    line(doc, "Почему у одной системы получилось несколько классов? Чем это отличается от описания "
              "системы одним названием?")
    for _ in range(4):
        line(doc, "________________________________________________________________________")
    doc.add_paragraph()

    line(doc, "Ответы на контрольные вопросы", bold=True)
    for i, q in enumerate(QUESTIONS, 1):
        line(doc, f"{i}. {q}")
        for _ in range(3):
            line(doc, "________________________________________________________________________")
        doc.add_paragraph()

    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
