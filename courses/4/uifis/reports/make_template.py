"""Generate the downloadable materials for UIFIS practices.

Run:  python make_template.py
Writes DOCX report templates, CSV tables and the practice 7 JSON fixture next
to this file for practices 1-9. Filenames are ASCII so the download links stay
safe in any browser.
"""

import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

HERE = Path(__file__).parent
OUT_DOCX = HERE / "UIFIS_practice-1_report-template.docx"
OUT_CSV = HERE / "UIFIS_practice-1_classification-table.csv"
OUT_DOCX2 = HERE / "UIFIS_practice-2_report-template.docx"
OUT_CSV2 = HERE / "UIFIS_practice-2_composition-table.csv"
OUT_DOCX3 = HERE / "UIFIS_practice-3_report-template.docx"
OUT_CSV3 = HERE / "UIFIS_practice-3_flow-registry.csv"
OUT_DOCX4 = HERE / "UIFIS_practice-4_report-template.docx"
OUT_CSV4 = HERE / "UIFIS_practice-4_decision-matrix.csv"
OUT_DOCX5 = HERE / "UIFIS_practice-5_report-template.docx"
OUT_CSV5 = HERE / "UIFIS_practice-5_configuration-matrix.csv"
OUT_DOCX6 = HERE / "UIFIS_practice-6_report-template.docx"
OUT_CSV6 = HERE / "UIFIS_practice-6_install-checklist.csv"
OUT_DOCX9 = HERE / "UIFIS_practice-9_report-template.docx"
OUT_CSV9 = HERE / "UIFIS_practice-9_service-checklist.csv"
QUESTIONS = [
    "Почему признаки классификации должны быть независимыми?",
    "Может ли одна система быть пакетной и оперативной одновременно?",
    "Чем масштаб отличается от сферы применения?",
]
QUESTIONS2 = [
    "Чем функциональная подсистема отличается от обеспечивающей?",
    "Почему функциональная подсистема не работает без обеспечивающих?",
    "К какому виду обеспечения относятся базы данных и почему?",
]
ROWS = [
    ("1", "Банковская АБС"),
    ("2", "1С:Бухгалтерия малого предприятия"),
    ("3", "Система диспетчеризации завода"),
    ("4", "Интернет-магазин"),
    ("5", "Медицинская картотека поликлиники"),
]
SYSTEMS2 = [
    "Электронный журнал колледжа",
    "Складской учёт",
    "Онлайн-запись к врачу",
]
QUESTIONS3 = [
    "Чем поток данных отличается от потока управления?",
    "Какие четыре элемента используются на диаграмме потоков данных?",
    "Почему у каждого процесса должен быть входной поток?",
    "Почему накопитель нельзя рисовать источником без процесса?",
]
QUESTIONS4 = [
    "Чем архитектура отличается от платформы?",
    "Почему в матрице выбора веса критериев должны давать единицу?",
    "Когда клиент-серверная архитектура лучше трёхзвенной?",
    "Почему схема размещения должна показывать границу доступа?",
    "Почему RAID не заменяет резервное копирование?",
]
CRITERIA4 = [
    ("Пользователи и нагрузка", "0,25"),
    ("Время ответа", "0,15"),
    ("Безопасность и доступ из дома", "0,20"),
    ("Частота обновления логики", "0,15"),
    ("Бюджет (5 значит дешевле)", "0,15"),
    ("Сопровождение", "0,10"),
]
QUESTIONS6 = [
    "Почему обновления устанавливают до настройки служб, а не после?",
    "Почему вход по SSH-ключу безопаснее пароля и что проверить до отключения пароля?",
    "Зачем приложению отдельный пользователь uifis, если уже есть student с правами sudo?",
    "Что проверить, если после смены статического адреса SSH перестал подключаться?",
    "Почему снимок виртуальной машины не заменяет резервную копию данных?",
]
CHECKS6 = [
    ("Имя узла", "hostnamectl", "Static hostname: uifis-srv01"),
    ("Часовой пояс", "timedatectl", "Asia/Yekaterinburg, синхронизация включена"),
    ("Сеть", "ip -br a", "адрес из задания, интерфейс UP"),
    ("Маршрут", "ip route", "маршрут по умолчанию есть"),
    ("SSH", "systemctl status ssh", "active (running)"),
    ("Вход по ключу", "ssh student@адрес", "вход без пароля"),
    ("Firewall", "sudo ufw status verbose", "активен, разрешён только SSH"),
    ("Службы", "systemctl --failed", "нет сбойных юнитов"),
    ("Диск", "df -h, df -i", "есть свободное место и inode"),
    ("Память", "free -h", "память видна, swap не перегружен"),
    ("Каталог", "ls -ld /srv/uifis", "владелец uifis, права 0750"),
    ("Снимок", "интерфейс гипервизора", "снимок after-base-setup существует"),
]
QUESTIONS5 = [
    "Почему для расчёта берут пиковую нагрузку, а не среднюю?",
    "Зачем нужен запас 30 процентов по памяти и диску?",
    "Почему кэш СУБД включают в расчёт памяти?",
    "Почему RAID не заменяет резервную копию?",
    "Когда отказоустойчивая конфигурация избыточна?",
]
CALC5 = [
    ("Процессор", "запросы/с × время на запрос / целевая загрузка", "", "", ""),
    ("Память", "(ОС + приложение + кэш) × 1,3", "", "", ""),
    ("Диск", "(данные + журналы + временные) × 1,3", "", "", ""),
    ("Сеть", "запросы/с × размер ответа × 8 × 2", "", "", ""),
]
CONFIG5 = [
    ("Процессор", "", "", "", ""),
    ("Память", "", "", "", ""),
    ("Диск", "", "", "", ""),
    ("RAID", "", "", "", ""),
    ("Сеть", "", "", "", ""),
    ("Питание", "", "", "", ""),
    ("Резервная копия", "", "", "", ""),
    ("Относительная стоимость", "", "", "", ""),
]
CHECKS5 = [
    ("Загрузка процессора", "не выше 65-70 процентов в пике", "", ""),
    ("Память", "запас не менее 30 процентов", "", ""),
    ("Диск", "запас не менее 30 процентов", "", ""),
    ("Сеть", "запас не менее 2 раз", "", ""),
    ("Диски", "RAID 1 или RAID 10 для системы", "", ""),
    ("Питание", "UPS для рабочего узла", "", ""),
    ("Резервная копия", "хранится отдельно", "", ""),
    ("Доступность", "соответствует RTO и RPO", "", ""),
]


def line(doc, text, bold=False, size=12, align=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    if align is not None:
        p.alignment = align
    return p


def slot(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(11)


def make_docx():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)

    line(doc, "Устройство и функционирование информационной системы (МДК.06.03)", bold=True, size=14,
         align=WD_ALIGN_PARAGRAPH.CENTER)
    line(doc, "Отчёт по практической работе №1", bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER)
    line(doc, "Тема: Классификация информационных систем", align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    line(doc, "ФИО студента: ______________________________________")
    line(doc, "Группа: ______________________  Дата: _______________")
    doc.add_paragraph()

    line(doc, "Цель работы", bold=True)
    line(doc, "________________________________________________________________________")
    line(doc, "________________________________________________________________________")
    doc.add_paragraph()

    line(doc, "Краткие теоретические сведения", bold=True)
    line(doc, "Перечислите четыре признака классификации ИС и приведите по 2 или 3 значения каждого:")
    line(doc, "1. Сфера применения: ____________________________________________________")
    line(doc, "2. Масштаб: _____________________________________________________________")
    line(doc, "3. Способ обработки: ____________________________________________________")
    line(doc, "4. Архитектура: _________________________________________________________")
    doc.add_paragraph()

    line(doc, "Таблица результатов классификации", bold=True)
    cols = ["№", "Информационная система", "Сфера", "Масштаб", "Способ обработки", "Архитектура",
            "Обоснование"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
    for n, name in ROWS:
        cells = t.add_row().cells
        cells[0].text = n
        cells[1].text = name
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
    doc.add_paragraph()

    line(doc, "Скриншоты", bold=True)
    slot(doc, "[ Место для скриншота 1: шапка отчёта с заполненными полями ]")
    doc.add_paragraph()
    slot(doc, "[ Место для скриншота 2: карточка «Банковская АБС» в классификаторе ]")
    doc.add_paragraph()

    line(doc, "Вывод", bold=True)
    line(doc, "Почему у одной системы получилось несколько классов? Чем это отличается от описания "
              "системы одним названием?")
    for _ in range(4):
        line(doc, "________________________________________________________________________")
    doc.add_paragraph()

    line(doc, "Ответы на контрольные вопросы", bold=True)
    for i, q in enumerate(QUESTIONS, 1):
        line(doc, f"{i}. {q}")
        for _ in range(3):
            line(doc, "________________________________________________________________________")
        doc.add_paragraph()

    doc.save(OUT_DOCX)
    print(f"wrote {OUT_DOCX}")


def make_csv():
    with OUT_CSV.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["№", "Информационная система", "Сфера", "Масштаб", "Способ обработки",
                    "Архитектура", "Обоснование"])
        for n, name in ROWS:
            w.writerow([n, name, "", "", "", "", ""])
    print(f"wrote {OUT_CSV}")


def make_docx2():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)

    line(doc, "Устройство и функционирование информационной системы (МДК.06.03)", bold=True, size=14,
         align=WD_ALIGN_PARAGRAPH.CENTER)
    line(doc, "Отчёт по практической работе №2", bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER)
    line(doc, "Тема: Построение схемы состава информационной системы", align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    line(doc, "ФИО студента: ______________________________________")
    line(doc, "Группа: ______________________  Дата: _______________")
    doc.add_paragraph()

    line(doc, "Цель работы", bold=True)
    line(doc, "________________________________________________________________________")
    line(doc, "________________________________________________________________________")
    doc.add_paragraph()

    line(doc, "Краткие теоретические сведения", bold=True)
    line(doc, "Перечислите пять видов обеспечения ИС и приведите пример каждого:")
    line(doc, "1. Техническое: _________________________________________________________")
    line(doc, "2. Программное: _________________________________________________________")
    line(doc, "3. Информационное: _____________________________________________________")
    line(doc, "4. Организационное: ____________________________________________________")
    line(doc, "5. Правовое: ___________________________________________________________")
    doc.add_paragraph()

    line(doc, "Схема состава информационной системы", bold=True)
    line(doc, "Нарисуйте схему состава выбранной системы: функциональные подсистемы, "
              "обеспечивающие подсистемы и связи между ними (какая функция на какое "
              "обеспечение опирается).")
    slot(doc, "[ Место для схемы состава: нарисуйте от руки, в графическом редакторе или вставьте снимок ]")
    doc.add_paragraph()

    line(doc, "Таблица состава", bold=True)
    cols = ["Часть", "Подсистема", "Что включает", "Опирается на"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
    for label in ("Функциональная",) * 3 + ("Обеспечивающая",) * 5:
        cells = t.add_row().cells
        cells[0].text = label
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
    doc.add_paragraph()

    line(doc, "Скриншоты", bold=True)
    slot(doc, "[ Место для скриншота 1: шапка отчёта с заполненными полями ]")
    doc.add_paragraph()
    slot(doc, "[ Место для скриншота 2: схема состава в рабочем виде ]")
    doc.add_paragraph()

    line(doc, "Вывод", bold=True)
    line(doc, "Почему функциональная подсистема не работает без обеспечивающих? Приведите "
              "пример из своей схемы.")
    for _ in range(4):
        line(doc, "________________________________________________________________________")
    doc.add_paragraph()

    line(doc, "Ответы на контрольные вопросы", bold=True)
    for i, q in enumerate(QUESTIONS2, 1):
        line(doc, f"{i}. {q}")
        for _ in range(3):
            line(doc, "________________________________________________________________________")
        doc.add_paragraph()

    doc.save(OUT_DOCX2)
    print(f"wrote {OUT_DOCX2}")


def make_csv2():
    with OUT_CSV2.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Часть", "Подсистема", "Что включает", "Опирается на"])
        for _ in range(3):
            w.writerow(["Функциональная", "", "", ""])
        for _ in range(5):
            w.writerow(["Обеспечивающая", "", "", ""])
    print(f"wrote {OUT_CSV2}")


def make_docx3():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)

    line(doc, "Устройство и функционирование информационной системы (МДК.06.03)", bold=True, size=14,
         align=WD_ALIGN_PARAGRAPH.CENTER)
    line(doc, "Отчёт по практической работе №3", bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER)
    line(doc, "Тема: Моделирование информационных потоков", align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    line(doc, "ФИО студента: ______________________________________")
    line(doc, "Группа: ______________________  Дата: _______________")
    doc.add_paragraph()

    line(doc, "Цель работы", bold=True)
    line(doc, "________________________________________________________________________")
    line(doc, "________________________________________________________________________")
    doc.add_paragraph()

    line(doc, "Краткие теоретические сведения", bold=True)
    line(doc, "Перечислите четыре элемента диаграммы потоков данных и их обозначения:")
    line(doc, "1. Внешняя сущность: ____________________________________________________")
    line(doc, "2. Процесс: _____________________________________________________________")
    line(doc, "3. Накопитель данных: ___________________________________________________")
    line(doc, "4. Поток данных: ________________________________________________________")
    doc.add_paragraph()

    line(doc, "Выбранная информационная система", bold=True)
    line(doc, "Название: _______________________________________________________________")
    doc.add_paragraph()

    line(doc, "Элементы системы", bold=True)
    cols = ["Тип элемента", "Название", "Что делает или хранит"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
    for label in ("Внешняя сущность",) * 2 + ("Процесс",) * 3 + ("Накопитель",) * 2:
        cells = t.add_row().cells
        cells[0].text = label
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
    doc.add_paragraph()

    line(doc, "Диаграмма потоков данных", bold=True)
    line(doc, "Нарисуйте диаграмму выбранной системы: внешние сущности, процессы, накопители "
              "и подписанные потоки данных. Если есть поток управления, покажите его отдельно.")
    slot(doc, "[ Место для диаграммы: нарисуйте от руки, в графическом редакторе или вставьте снимок ]")
    doc.add_paragraph()

    line(doc, "Реестр потоков", bold=True)
    cols = ["№", "Состав данных", "Источник", "Приёмник", "Тип"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
    for n in range(1, 9):
        cells = t.add_row().cells
        cells[0].text = str(n)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
    doc.add_paragraph()

    line(doc, "Скриншоты", bold=True)
    slot(doc, "[ Место для скриншота 1: шапка отчёта с заполненными полями ]")
    doc.add_paragraph()
    slot(doc, "[ Место для скриншота 2: диаграмма потоков данных в рабочем виде ]")
    doc.add_paragraph()

    line(doc, "Вывод", bold=True)
    line(doc, "Какие процессы в вашей системе преобразуют данные и где они хранятся? Приведите "
              "пример потока, который было бы легко потерять.")
    for _ in range(4):
        line(doc, "________________________________________________________________________")
    doc.add_paragraph()

    line(doc, "Ответы на контрольные вопросы", bold=True)
    for i, q in enumerate(QUESTIONS3, 1):
        line(doc, f"{i}. {q}")
        for _ in range(3):
            line(doc, "________________________________________________________________________")
        doc.add_paragraph()

    doc.save(OUT_DOCX3)
    print(f"wrote {OUT_DOCX3}")


def make_csv3():
    with OUT_CSV3.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["№", "Состав данных", "Источник", "Приёмник", "Тип"])
        for n in range(1, 9):
            w.writerow([n, "", "", "", ""])
    print(f"wrote {OUT_CSV3}")


def make_docx4():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)

    line(doc, "Устройство и функционирование информационной системы (МДК.06.03)", bold=True, size=14,
         align=WD_ALIGN_PARAGRAPH.CENTER)
    line(doc, "Отчёт по практической работе №4", bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER)
    line(doc, "Тема: Проектирование архитектуры информационной системы",
         align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    line(doc, "ФИО студента: ______________________________________")
    line(doc, "Группа: ______________________  Дата: _______________")
    doc.add_paragraph()

    line(doc, "Цель работы", bold=True)
    line(doc, "________________________________________________________________________")
    line(doc, "________________________________________________________________________")
    doc.add_paragraph()

    line(doc, "Краткие теоретические сведения", bold=True)
    line(doc, "Опишите три слоя и три базовых варианта размещения:")
    line(doc, "1. Представление: _______________________________________________________")
    line(doc, "2. Логика: ______________________________________________________________")
    line(doc, "3. Данные: ______________________________________________________________")
    line(doc, "4. Файл-сервер: _________________________________________________________")
    line(doc, "5. Клиент-сервер: _______________________________________________________")
    line(doc, "6. Трёхзвенная: _________________________________________________________")
    doc.add_paragraph()

    line(doc, "Выбранная система и профиль нагрузки", bold=True)
    line(doc, "Название: _______________________________________________________________")
    line(doc, "Пользователи: ___________________________________________________________")
    line(doc, "Одновременные сессии: __________________________________________________")
    line(doc, "Данные и доступ: ________________________________________________________")
    doc.add_paragraph()

    line(doc, "Требования", bold=True)
    cols = ["Тип", "Номер", "Требование", "Значение или проверка"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
    for kind, prefix, count in (("Функциональное", "ФТ", 4), ("Нефункциональное", "НФТ", 5)):
        for n in range(1, count + 1):
            cells = t.add_row().cells
            cells[0].text = kind
            cells[1].text = f"{prefix}-{n}"
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    doc.add_paragraph()

    line(doc, "Варианты архитектуры", bold=True)
    cols = ["Вариант", "Представление", "Логика", "Данные", "Что идёт по сети"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
    for name in ("Файл-сервер", "Клиент-сервер", "Трёхзвенная"):
        cells = t.add_row().cells
        cells[0].text = name
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    doc.add_paragraph()

    line(doc, "Матрица выбора", bold=True)
    line(doc, "Оценки от 1 до 5. Для бюджета 5 значит дешевле. Сумма весов равна 1,00.")
    cols = ["Критерий", "Вес", "Файл-сервер", "Клиент-сервер", "Трёхзвенная"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
    for criterion, weight in CRITERIA4:
        cells = t.add_row().cells
        cells[0].text = criterion
        cells[1].text = weight
    cells = t.add_row().cells
    cells[0].text = "Взвешенная оценка"
    cells[1].text = "1,00"
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    doc.add_paragraph()

    line(doc, "Решение", bold=True)
    line(doc, "Выбранный вариант: ______________________________________________________")
    line(doc, "________________________________________________________________________")
    line(doc, "Какие требования он закрывает: __________________________________________")
    line(doc, "________________________________________________________________________")
    line(doc, "Чем придётся заплатить: _________________________________________________")
    line(doc, "________________________________________________________________________")
    doc.add_paragraph()

    line(doc, "Схема размещения", bold=True)
    line(doc, "Покажите узлы, слои на них, сетевые связи, подписи и границу доступа.")
    slot(doc, "[ Место для схемы: нарисуйте от руки, в графическом редакторе или вставьте снимок ]")
    doc.add_paragraph()

    line(doc, "Таблица узлов", bold=True)
    cols = ["Узел", "Роль", "Программное обеспечение", "Сеть", "Резерв"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
    for _ in range(5):
        t.add_row()
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    doc.add_paragraph()

    line(doc, "Риски решения", bold=True)
    cols = ["Риск", "Вероятность", "Влияние", "Мера снижения"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
    for _ in range(4):
        t.add_row()
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    doc.add_paragraph()

    line(doc, "Скриншоты", bold=True)
    slot(doc, "[ Место для скриншота 1: шапка отчёта и профиль системы ]")
    doc.add_paragraph()
    slot(doc, "[ Место для скриншота 2: заполненная матрица выбора ]")
    doc.add_paragraph()
    slot(doc, "[ Место для скриншота 3: схема размещения с подписями ]")
    doc.add_paragraph()

    line(doc, "Вывод", bold=True)
    line(doc, "Какое требование было самым трудным? Почему выбранный вариант его закрывает? "
              "Что изменилось бы при другом бюджете?")
    for _ in range(6):
        line(doc, "________________________________________________________________________")
    doc.add_paragraph()

    line(doc, "Ответы на контрольные вопросы", bold=True)
    for i, q in enumerate(QUESTIONS4, 1):
        line(doc, f"{i}. {q}")
        for _ in range(3):
            line(doc, "________________________________________________________________________")
        doc.add_paragraph()

    doc.save(OUT_DOCX4)
    print(f"wrote {OUT_DOCX4}")


def make_csv4():
    with OUT_CSV4.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Критерий", "Вес", "Файл-сервер", "Клиент-сервер", "Трёхзвенная",
                    "Комментарий"])
        for criterion, weight in CRITERIA4:
            w.writerow([criterion, weight, "", "", "", ""])
        w.writerow(["Взвешенная оценка", "1,00", "", "", "", ""])
    print(f"wrote {OUT_CSV4}")


def make_docx6():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)

    line(doc, "Устройство и функционирование информационной системы (МДК.06.03)", bold=True, size=14,
         align=WD_ALIGN_PARAGRAPH.CENTER)
    line(doc, "Отчёт по практической работе №6", bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER)
    line(doc, "Тема: Установка и базовая настройка серверной операционной системы",
         align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    line(doc, "ФИО студента: ______________________________________")
    line(doc, "Группа: ______________________  Дата: _______________")
    doc.add_paragraph()

    line(doc, "Цель работы", bold=True)
    line(doc, "________________________________________________________________________")
    line(doc, "________________________________________________________________________")
    doc.add_paragraph()

    line(doc, "Параметры виртуальной машины", bold=True)
    cols = ["Параметр", "Значение"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
    for name in ("Имя узла", "Гипервизор", "Память", "Процессоры", "Диск", "Сеть", "Версия ОС"):
        cells = t.add_row().cells
        cells[0].text = name
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
    doc.add_paragraph()

    line(doc, "Журнал установки", bold=True)
    cols = ["Этап", "Что сделано", "Команда или наблюдение", "Результат"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
    for n in range(1, 8):
        cells = t.add_row().cells
        cells[0].text = str(n)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    doc.add_paragraph()

    line(doc, "Базовая настройка", bold=True)
    cols = ["Настройка", "Команда", "Фактический результат"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
    for name in ("Имя узла", "Обновления", "Часовой пояс", "Статический адрес",
                 "SSH и ключи", "Firewall", "Служебный пользователь", "Снимок"):
        cells = t.add_row().cells
        cells[0].text = name
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    doc.add_paragraph()

    line(doc, "Чек-лист приёмки", bold=True)
    cols = ["Проверка", "Команда", "Ожидаемый результат", "Факт", "Отметка"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
    for name, command, expected in CHECKS6:
        cells = t.add_row().cells
        cells[0].text = name
        cells[1].text = command
        cells[2].text = expected
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8)
    doc.add_paragraph()

    line(doc, "Скриншоты", bold=True)
    for caption in (
        "первый экран установщика",
        "вход после первой загрузки",
        "вывод ip -br a",
        "вывод systemctl --failed",
        "вывод sudo ufw status verbose",
        "успешный вход по SSH",
    ):
        slot(doc, f"[ Место для скриншота: {caption} ]")
        doc.add_paragraph()

    line(doc, "Трудности и их решение", bold=True)
    line(doc, "Какая настройка вызвала трудности? Как вы определили причину и чем проверили решение?")
    for _ in range(6):
        line(doc, "________________________________________________________________________")
    doc.add_paragraph()

    line(doc, "Вывод", bold=True)
    line(doc, "Что доказало, что узел готов к работе? Какие проверки вы считаете обязательными "
              "для любого сервера?")
    for _ in range(5):
        line(doc, "________________________________________________________________________")
    doc.add_paragraph()

    line(doc, "Ответы на контрольные вопросы", bold=True)
    for i, q in enumerate(QUESTIONS6, 1):
        line(doc, f"{i}. {q}")
        for _ in range(3):
            line(doc, "________________________________________________________________________")
        doc.add_paragraph()

    doc.save(OUT_DOCX6)
    print(f"wrote {OUT_DOCX6}")


def make_csv6():
    with OUT_CSV6.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Проверка", "Команда", "Ожидаемый результат", "Фактический результат",
                    "Отметка"])
        for name, command, expected in CHECKS6:
            w.writerow([name, command, expected, "", ""])
        w.writerow(["Итог", "принято или не принято", "", "", ""])
    print(f"wrote {OUT_CSV6}")


def make_docx5():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)

    line(doc, "Устройство и функционирование информационной системы (МДК.06.03)", bold=True, size=14,
         align=WD_ALIGN_PARAGRAPH.CENTER)
    line(doc, "Отчёт по практической работе №5", bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER)
    line(doc, "Тема: Подбор аппаратной конфигурации информационной системы",
         align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    line(doc, "ФИО студента: ______________________________________")
    line(doc, "Группа: ______________________  Дата: _______________")
    doc.add_paragraph()

    line(doc, "Цель работы", bold=True)
    line(doc, "________________________________________________________________________")
    line(doc, "________________________________________________________________________")
    doc.add_paragraph()

    line(doc, "Профиль системы", bold=True)
    cols = ["Параметр", "Значение"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
    for name in ("Система", "Пользователи", "Пиковые сессии", "Запросы в секунду",
                 "Средний размер ответа", "Объём данных и рост", "RTO", "RPO", "Бюджет"):
        cells = t.add_row().cells
        cells[0].text = name
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    doc.add_paragraph()

    line(doc, "Расчёт ресурсов", bold=True)
    line(doc, "Покажите формулу, подстановку и итог. Округляйте вверх.")
    cols = ["Ресурс", "Формула", "Расчёт", "Значение", "Запас"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
    for name, formula, _, _, _ in CALC5:
        cells = t.add_row().cells
        cells[0].text = name
        cells[1].text = formula
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8)
    doc.add_paragraph()

    line(doc, "Конфигурации", bold=True)
    cols = ["Компонент", "Минимальная", "Рекомендуемая", "Отказоустойчивая", "Комментарий"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
    for name, _, _, _, _ in CONFIG5:
        cells = t.add_row().cells
        cells[0].text = name
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8)
    doc.add_paragraph()

    line(doc, "Проверка расчёта", bold=True)
    cols = ["Что проверить", "Норма", "Факт", "Вывод"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
    for name, norm, _, _ in CHECKS5:
        cells = t.add_row().cells
        cells[0].text = name
        cells[1].text = norm
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8)
    doc.add_paragraph()

    line(doc, "Выбор конфигурации", bold=True)
    line(doc, "Выбранная конфигурация: ________________________________________________")
    line(doc, "Какие требования она закрывает: ________________________________________")
    line(doc, "________________________________________________________________________")
    line(doc, "Чем придётся заплатить: _________________________________________________")
    line(doc, "________________________________________________________________________")
    line(doc, "Почему не выбраны два других варианта: __________________________________")
    line(doc, "________________________________________________________________________")
    doc.add_paragraph()

    line(doc, "Риски выбранного решения", bold=True)
    cols = ["Риск", "Вероятность", "Влияние", "Мера снижения"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
    for _ in range(4):
        t.add_row()
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    doc.add_paragraph()

    line(doc, "Скриншоты", bold=True)
    for caption in (
        "профиль системы и требования",
        "расчётная таблица с формулами",
        "итоговая матрица конфигураций",
    ):
        slot(doc, f"[ Место для скриншота: {caption} ]")
        doc.add_paragraph()

    line(doc, "Вывод", bold=True)
    line(doc, "Какое требование было самым трудным? Почему выбранная конфигурация достаточна "
              "и что изменилось бы при другом бюджете?")
    for _ in range(5):
        line(doc, "________________________________________________________________________")
    doc.add_paragraph()

    line(doc, "Ответы на контрольные вопросы", bold=True)
    for i, q in enumerate(QUESTIONS5, 1):
        line(doc, f"{i}. {q}")
        for _ in range(3):
            line(doc, "________________________________________________________________________")
        doc.add_paragraph()

    doc.save(OUT_DOCX5)
    print(f"wrote {OUT_DOCX5}")


def make_csv5():
    with OUT_CSV5.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Компонент", "Минимальная", "Рекомендуемая", "Отказоустойчивая",
                    "Комментарий"])
        for name, _, _, _, _ in CONFIG5:
            w.writerow([name, "", "", "", ""])
        w.writerow(["Итог", "выбрано или не выбрано", "", "", ""])
    print(f"wrote {OUT_CSV5}")


OUT_DOCX7 = HERE / "UIFIS_practice-7_report-template.docx"
OUT_JSON7 = HERE / "UIFIS_practice-7_students.json"
OUT_DOCX8 = HERE / "UIFIS_practice-8_report-template.docx"
OUT_CSV8 = HERE / "UIFIS_practice-8_diagnostics.csv"


def set_table_font(table, size=9):
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(size)


def make_json7():
    students = [
        {
            "full_name": "Иванова Мария Сергеевна",
            "email": "ivanova@example.ru",
            "group_code": "ИС-41",
        },
        {
            "full_name": "Крылов Артём Игоревич",
            "email": "krylov@example.ru",
            "group_code": "ИС-42",
        },
        {
            "full_name": "Мельникова Ольга Павловна",
            "email": "melnikova@example.ru",
            "group_code": "ИС-41",
        },
    ]
    with OUT_JSON7.open("w", encoding="utf-8", newline="\n") as f:
        json.dump(students, f, ensure_ascii=False, indent=2)
        f.write("\n")
    print(f"wrote {OUT_JSON7}")


def make_docx7():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)

    line(doc, "Устройство и функционирование информационной системы (МДК.06.03)",
         bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
    line(doc, "Отчёт по практической работе 7", bold=True, size=13,
         align=WD_ALIGN_PARAGRAPH.CENTER)
    line(doc, "Тема: Создание базы данных информационной системы",
         align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    line(doc, "ФИО студента: ______________________________________")
    line(doc, "Группа: ______________________  Дата: _______________")
    doc.add_paragraph()

    line(doc, "Цель работы", bold=True)
    line(doc, "Создать базу данных учебного журнала в SQLite, задать ограничения "
              "целостности, наполнить таблицы данными, импортировать JSON и проверить "
              "выборки с соединением и агрегацией.")
    doc.add_paragraph()

    line(doc, "Схема базы данных", bold=True)
    cols = ["Таблица", "Назначение", "Ключи и ограничения", "Строк"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
    for row in (
        ("students", "Студенты", "PK id, UNIQUE email, CHECK status", "6"),
        ("courses", "Курсы", "PK id, UNIQUE code, CHECK hours", "4"),
        ("enrollments", "Записи на курсы", "PK id, FK student, FK course, UNIQUE пара", "8"),
        ("grades", "Оценки", "PK id, FK enrollment, UNIQUE enrollment, CHECK score", "7"),
    ):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    set_table_font(t, 9)
    doc.add_paragraph()

    line(doc, "ER-модель в третьей нормальной форме", bold=True)
    slot(doc, "[ Вставьте схему из draw.io: students, courses, enrollments, grades ]")
    doc.add_paragraph()
    line(doc, "Объясните, почему оценка хранится в отдельной таблице grades, а не в enrollments:")
    for _ in range(4):
        line(doc, "________________________________________________________________________")
    doc.add_paragraph()

    line(doc, "Проверка ограничений", bold=True)
    cols = ["Проверка", "Действие", "Ожидаемый результат", "Факт", "Вывод"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
    for name, action, expected in (
        ("Оценка 101", "INSERT grade score = 101", "CHECK отклоняет строку"),
        ("Повтор email", "INSERT student с тем же email", "UNIQUE отклоняет строку"),
        ("Удаление курса", "DELETE course с записями", "RESTRICT отклоняет удаление"),
        ("Удаление студента", "DELETE student с оценками", "CASCADE удаляет записи и оценки"),
        ("Повтор записи", "Повторная запись на тот же курс", "UNIQUE отклоняет строку"),
    ):
        cells = t.add_row().cells
        cells[0].text = name
        cells[1].text = action
        cells[2].text = expected
    set_table_font(t, 8)
    doc.add_paragraph()

    line(doc, "Результаты SQL-запросов", bold=True)
    cols = ["Запрос", "Что проверяет", "Результат", "Скриншот"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
    for name, purpose in (
        ("Запрос 1", "Число курсов у каждого студента"),
        ("Запрос 2", "Средний балл по каждому курсу"),
        ("Запрос 3", "Студенты без записей на курсы"),
        ("Запрос 4", "Три лучших студента по среднему баллу"),
        ("Запрос 5", "Курсы без выставленных оценок"),
    ):
        cells = t.add_row().cells
        cells[0].text = name
        cells[1].text = purpose
    set_table_font(t, 9)
    doc.add_paragraph()

    line(doc, "Импорт JSON", bold=True)
    cols = ["Показатель", "Значение"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
    for name in ("Файл", "Строк в файле", "Импортировано", "Отклонено", "Итоговая проверка"):
        t.add_row().cells[0].text = name
    set_table_font(t, 9)
    doc.add_paragraph()

    line(doc, "Скриншоты", bold=True)
    for caption in (
        "Схема SQLite и список таблиц",
        "Ошибка CHECK при score = 101",
        "Импорт JSON и сообщение об успехе",
        "Результат запроса 2 со средним баллом",
    ):
        slot(doc, f"[ Место для скриншота: {caption} ]")
        doc.add_paragraph()

    line(doc, "Вывод", bold=True)
    line(doc, "Какое ограничение оказалось самым важным? Что изменилось бы, если бы "
              "внешние ключи были выключены? Где в экзаменационном задании пригодится "
              "этот скрипт?")
    for _ in range(6):
        line(doc, "________________________________________________________________________")
    doc.add_paragraph()

    line(doc, "Ответы на контрольные вопросы", bold=True)
    for i, q in enumerate((
        "Почему таблица grades отделена от enrollments?",
        "Что произойдёт при удалении студента, у которого есть оценки?",
        "Зачем нужен PRAGMA foreign_keys = ON?",
        "Чем UNIQUE отличается от PRIMARY KEY?",
        "Почему импорт JSON выполняют в одной транзакции?",
    ), 1):
        line(doc, f"{i}. {q}")
        for _ in range(3):
            line(doc, "________________________________________________________________________")
        doc.add_paragraph()

    doc.save(OUT_DOCX7)
    print(f"wrote {OUT_DOCX7}")


def make_docx8():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)

    line(doc, "Устройство и функционирование информационной системы (МДК.06.03)",
         bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
    line(doc, "Отчёт по практической работе 8", bold=True, size=13,
         align=WD_ALIGN_PARAGRAPH.CENTER)
    line(doc, "Тема: Настройка сети и диагностика соединений",
         align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    line(doc, "ФИО студента: ______________________________________")
    line(doc, "Группа: ______________________  Дата: _______________")
    doc.add_paragraph()

    line(doc, "Цель работы", bold=True)
    line(doc, "Связать две учебные виртуальные машины внутренней сетью, настроить "
              "адреса и проверить доступность сервиса, а затем найти и исправить "
              "пять типовых неисправностей.")
    doc.add_paragraph()

    line(doc, "Схема сети", bold=True)
    cols = ["Узел", "Роль", "Адаптер 1", "Адаптер 2", "Внутренний адрес"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
    for row in (
        ("uifis-srv01", "Сервер", "NAT, DHCP", "uifis-net", "10.10.0.10/24"),
        ("uifis-cli01", "Клиент", "NAT, DHCP", "uifis-net", "10.10.0.20/24"),
    ):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    set_table_font(t, 9)
    doc.add_paragraph()
    slot(doc, "[ Вставьте схему сети: две ВМ, NAT и внутренняя сеть uifis-net ]")
    doc.add_paragraph()

    line(doc, "Параметры внутренней сети", bold=True)
    cols = ["Параметр", "Значение"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
    for name in ("Имя сети VirtualBox", "Адрес сервера", "Адрес клиента", "Маска",
                 "Шлюз по умолчанию", "DNS"):
        t.add_row().cells[0].text = name
    set_table_font(t, 10)
    doc.add_paragraph()

    line(doc, "Проверки", bold=True)
    cols = ["Проверка", "Команда", "Ожидаемый результат", "Факт", "Вывод"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
    checks = (
        ("Адреса", "ip -br a", "10.10.0.10/24 и 10.10.0.20/24, state UP"),
        ("Маршруты", "ip route", "Маршрут 10.10.0.0/24 и default через NAT"),
        ("Связность", "ping -c 3 10.10.0.20", "3 пакета, потерь 0%"),
        ("Служба", "ss -tlnp | grep 8080", "служба слушает 10.10.0.10:8080"),
        ("HTTP", "curl -I http://10.10.0.10:8080/", "HTTP/1.0 200 OK"),
        ("SSH", "ssh student@10.10.0.10", "вход по паролю"),
        ("Firewall", "sudo ufw status", "разрешён TCP 8080 из 10.10.0.0/24"),
        ("Внешний сайт", "curl -I https://example.com", "HTTP 200 через NAT"),
    )
    for row in checks:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    set_table_font(t, 8)
    doc.add_paragraph()

    line(doc, "Неисправности и исправления", bold=True)
    cols = ["Неисправность", "Признак", "Проверка", "Исправление"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
    faults = (
        ("Неверная подсеть", "ping не проходит", "ip -br a", "вернуть 10.10.0.20/24"),
        ("Служба на 127.0.0.1", "connection refused", "ss -tlnp", "слушать 10.10.0.10"),
        ("Закрыт порт", "curl зависает", "sudo ufw status", "разрешить 8080"),
        ("Нет default route", "нет выхода в интернет", "ip route", "вернуть NAT-маршрут"),
        ("Неверный DNS", "имя не разрешается", "getent hosts example.com", "вернуть DNS NAT"),
    )
    for row in faults:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    set_table_font(t, 8)
    doc.add_paragraph()

    line(doc, "Скриншоты", bold=True)
    for caption in (
        "Настройки двух адаптеров в VirtualBox",
        "Вывод ip -br a на сервере и клиенте",
        "Успешный curl с клиента",
        "Вывод ss -tlnp",
        "Исправленная неисправность",
    ):
        slot(doc, f"[ Место для скриншота: {caption} ]")
        doc.add_paragraph()

    line(doc, "Вывод", bold=True)
    line(doc, "Какая неисправность была самой трудной? Как вы отличили ошибку маршрута "
              "от ошибки firewall? Почему внутренняя сеть не должна иметь свой шлюз?")
    for _ in range(6):
        line(doc, "________________________________________________________________________")
    doc.add_paragraph()

    line(doc, "Ответы на контрольные вопросы", bold=True)
    for i, q in enumerate((
        "Зачем нужны две сети: NAT и internal?",
        "Почему на внутреннем интерфейсе не задают default gateway?",
        "Что показывает команда ss -tlnp?",
        "Почему сервис на 127.0.0.1 недоступен с другой виртуальной машины?",
        "Как отличить ошибку маршрута от ошибки firewall?",
    ), 1):
        line(doc, f"{i}. {q}")
        for _ in range(3):
            line(doc, "________________________________________________________________________")
        doc.add_paragraph()

    doc.save(OUT_DOCX8)
    print(f"wrote {OUT_DOCX8}")


def make_csv8():
    checks = (
        ("Адреса", "ip -br a", "10.10.0.10/24 и 10.10.0.20/24, state UP"),
        ("Маршруты", "ip route", "10.10.0.0/24 и default через NAT"),
        ("Связность", "ping -c 3 10.10.0.20", "3 пакета, потерь 0%"),
        ("Служба", "ss -tlnp | grep 8080", "10.10.0.10:8080"),
        ("HTTP", "curl -I http://10.10.0.10:8080/", "HTTP/1.0 200 OK"),
        ("SSH", "ssh student@10.10.0.10", "вход по паролю"),
        ("Firewall", "sudo ufw status", "TCP 8080 разрешён из 10.10.0.0/24"),
        ("Внешний сайт", "curl -I https://example.com", "HTTP 200 через NAT"),
    )
    with OUT_CSV8.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Проверка", "Команда", "Ожидаемый результат", "Фактический результат", "Вывод"])
        for name, command, expected in checks:
            w.writerow([name, command, expected, "", ""])
        w.writerow(["Итог", "все проверки", "8 из 8", "", ""])
    print(f"wrote {OUT_CSV8}")


def make_docx9():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)

    line(doc, "Устройство и функционирование информационной системы (МДК.06.03)",
         bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
    line(doc, "Отчёт по практической работе 9", bold=True, size=13,
         align=WD_ALIGN_PARAGRAPH.CENTER)
    line(doc, "Тема: Развёртывание базовой сетевой службы",
         align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    line(doc, "ФИО студента: ______________________________________")
    line(doc, "Группа: ______________________  Дата: _______________")
    doc.add_paragraph()

    line(doc, "Цель работы", bold=True)
    line(doc, "Развернуть учебную HTTP-службу uifis-web в Ubuntu Server, оформить "
              "запуск через systemd, настроить отдельного пользователя, firewall и "
              "проброс порта, а затем проверить отказ и восстановление службы.")
    doc.add_paragraph()

    line(doc, "Схема стенда", bold=True)
    cols = ["Элемент", "Значение"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
    for name, value in (
        ("Рабочий компьютер", "Windows, PowerShell, curl.exe"),
        ("Виртуальная машина", "uifis-srv01, Ubuntu Server"),
        ("Сеть VirtualBox", "NAT, адрес гостя 10.0.2.15"),
        ("SSH", "хост 127.0.0.1:2222 → гость 22/tcp"),
        ("Служба", "uifis-web, HTTP, 0.0.0.0:8080"),
        ("Проброс порта", "хост 127.0.0.1:8080 → гость 8080/tcp"),
        ("Пользователь службы", "uifis, без интерактивного входа"),
        ("Каталог службы", "/srv/uifis-web"),
    ):
        cells = t.add_row().cells
        cells[0].text = name
        cells[1].text = value
    set_table_font(t, 9)
    doc.add_paragraph()
    slot(doc, "[ Вставьте схему стенда: хост, VirtualBox NAT, проброс порта, служба ]")
    doc.add_paragraph()

    line(doc, "Контракт службы", bold=True)
    cols = ["Метод", "Путь", "Код", "Тип ответа", "Назначение", "Факт"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
    for row in (
        ("GET", "/", "200", "text/html", "Главная страница службы", ""),
        ("GET", "/health.json", "200", "application/json", "Состояние и время работы", ""),
        ("GET", "/api/services", "200", "application/json", "Каталог сетевых служб", ""),
        ("GET", "/missing", "404", "application/json", "Ошибка для неизвестного пути", ""),
        ("HEAD", "/health.json", "200", "без тела", "Проверка заголовков", ""),
    ):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    set_table_font(t, 8)
    doc.add_paragraph()

    line(doc, "Параметры развёртывания", bold=True)
    cols = ["Параметр", "Значение", "Факт"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
    for name in ("Имя службы", "Пользователь", "Рабочая директория", "Файл приложения",
                 "Адрес прослушивания", "Порт", "Файл unit", "Правило firewall",
                 "Проброс порта"):
        cells = t.add_row().cells
        cells[0].text = name
    set_table_font(t, 9)
    doc.add_paragraph()

    line(doc, "Проверки", bold=True)
    cols = ["Проверка", "Команда", "Ожидаемый результат", "Факт", "Вывод"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
    checks = (
        ("Каталог", "ls -ld /srv/uifis-web", "владелец uifis, права 0750"),
        ("Файл", "ls -l /srv/uifis-web/app.py", "владелец uifis, права 0640"),
        ("Синтаксис", "sudo -u uifis python3 -m py_compile /srv/uifis-web/app.py", "без ошибок"),
        ("Служба", "systemctl status uifis-web", "active (running)"),
        ("Автозапуск", "systemctl is-enabled uifis-web", "enabled"),
        ("Порт", "ss -tulpn | grep 8080", "0.0.0.0:8080"),
        ("Главная", "curl -i http://127.0.0.1:8080/", "200 OK, HTML"),
        ("Здоровье", "curl -i http://127.0.0.1:8080/health.json", "200 OK, status ok"),
        ("Каталог API", "curl -i http://127.0.0.1:8080/api/services", "count равно 4"),
        ("Ошибка", "curl -i http://127.0.0.1:8080/missing", "404 Not Found"),
        ("Журнал", "journalctl -u uifis-web -n 30 --no-pager", "строки запуска и запросов"),
        ("Firewall", "sudo ufw status verbose", "8080/tcp из 10.0.2.0/24"),
        ("Проброс", "VirtualBox → Проброс портов", "TCP 8080 на 8080"),
        ("С хоста", "curl.exe -i http://127.0.0.1:8080/health.json", "JSON получен"),
        ("Отказ", "systemctl stop uifis-web", "служба не отвечает"),
        ("Восстановление", "systemctl start uifis-web", "снова 200 OK"),
    )
    for row in checks:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    set_table_font(t, 8)
    doc.add_paragraph()

    line(doc, "Отказы и восстановление", bold=True)
    cols = ["Отказ", "Симптом", "Проверка", "Причина", "Восстановление"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        t.rows[0].cells[i].text = c
    for name in ("Служба остановлена", "Служба слушает 127.0.0.1",
                 "Порт закрыт firewall", "Проброс ведёт на 8081",
                 "Файл app.py недоступен"):
        cells = t.add_row().cells
        cells[0].text = name
    set_table_font(t, 8)
    doc.add_paragraph()

    line(doc, "Скриншоты", bold=True)
    for caption in (
        "Вывод systemctl status uifis-web",
        "Вывод ss -tulpn с адресом 0.0.0.0:8080",
        "Ответы /, /health.json и /api/services из гостевой системы",
        "Ответ curl.exe с рабочего компьютера",
        "Правило проброса портов в VirtualBox",
        "Один отказ и его восстановление",
    ):
        slot(doc, f"[ Место для скриншота: {caption} ]")
        doc.add_paragraph()

    line(doc, "Вывод", bold=True)
    line(doc, "Какая зависимость оказалась критичной для доступа к службе? Почему "
              "локальная проверка не заменяет проверку с хоста? Чем systemd удобнее "
              "ручного запуска?")
    for _ in range(6):
        line(doc, "________________________________________________________________________")
    doc.add_paragraph()

    line(doc, "Ответы на контрольные вопросы", bold=True)
    for i, q in enumerate((
        "Почему ручной запуск python3 app.py не считается полноценным развёртыванием?",
        "Зачем службе отдельный пользователь uifis и права 0640 на файл?",
        "Что проверять, если служба активна, но с хоста не открывается?",
        "Почему недостаточно открыть порт в приложении, не проверив firewall?",
        "Чем 0.0.0.0:8080 отличается от 127.0.0.1:8080?",
        "Почему systemctl status не заменяет проверку /health.json?",
        "Чем проброс порта в VirtualBox удобнее изменения настроек хоста?",
    ), 1):
        line(doc, f"{i}. {q}")
        for _ in range(3):
            line(doc, "________________________________________________________________________")
        doc.add_paragraph()

    doc.save(OUT_DOCX9)
    print(f"wrote {OUT_DOCX9}")


def make_csv9():
    checks = (
        ("Каталог", "ls -ld /srv/uifis-web", "владелец uifis, права 0750"),
        ("Файл", "ls -l /srv/uifis-web/app.py", "владелец uifis, права 0640"),
        ("Синтаксис", "sudo -u uifis python3 -m py_compile /srv/uifis-web/app.py", "без ошибок"),
        ("Служба", "systemctl status uifis-web", "active (running)"),
        ("Автозапуск", "systemctl is-enabled uifis-web", "enabled"),
        ("Порт", "ss -tulpn | grep 8080", "0.0.0.0:8080"),
        ("Главная", "curl -i http://127.0.0.1:8080/", "200 OK, HTML"),
        ("Здоровье", "curl -i http://127.0.0.1:8080/health.json", "200 OK, status ok"),
        ("Каталог API", "curl -i http://127.0.0.1:8080/api/services", "count равно 4"),
        ("Ошибка", "curl -i http://127.0.0.1:8080/missing", "404 Not Found"),
        ("Журнал", "journalctl -u uifis-web -n 30 --no-pager", "строки запуска и запросов"),
        ("Firewall", "sudo ufw status verbose", "8080/tcp из 10.0.2.0/24"),
        ("Проброс", "VirtualBox → Проброс портов", "TCP 8080 на 8080"),
        ("С хоста", "curl.exe -i http://127.0.0.1:8080/health.json", "JSON получен"),
        ("Отказ", "systemctl stop uifis-web", "служба не отвечает"),
        ("Восстановление", "systemctl start uifis-web", "снова 200 OK"),
    )
    with OUT_CSV9.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Проверка", "Команда", "Ожидаемый результат", "Фактический результат", "Вывод"])
        for name, command, expected in checks:
            w.writerow([name, command, expected, "", ""])
        w.writerow(["Итог", "все проверки", "16 из 16", "", ""])
    print(f"wrote {OUT_CSV9}")


if __name__ == "__main__":
    make_docx()
    make_csv()
    make_docx2()
    make_csv2()
    make_docx3()
    make_csv3()
    make_docx4()
    make_csv4()
    make_docx5()
    make_csv5()
    make_docx6()
    make_csv6()
    make_json7()
    make_docx7()
    make_docx8()
    make_csv8()
    make_docx9()
    make_csv9()
